##Bulk RNA Seq Data Analysis Workflow by Md Jahid Hasan Jone##

# Summarize the fastp JSON reports (run after 1_Bulk_RNASeq_fastp.sh).
#
# Reads every .json file in 5_Results/1_fastp/ and writes one Excel table,
# one row per sample:
#
#     5_Results/1_fastp/fastp_summary.xlsx
#
# Columns: read counts, bases, Q20/Q30, mean read length and GC content
# before and after filtering, filtering results, duplication rate, insert
# size peak and adapter-trimmed reads. Read2 columns are empty for single-end data.
#
# Needs: pip install openpyxl

import os
import re
import sys
import glob
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# =====================================================================
# PATHS  <- EDIT THESE (replace /.../.../ with your path)
# =====================================================================
json_dir = "/.../.../Bulk_RNA_seq/5_Results/1_fastp"
output_file = "/.../.../Bulk_RNA_seq/5_Results/1_fastp/fastp_summary.xlsx"


def safe_get(d, *keys, default=""):
    """Safely walk nested dict keys, returning default if missing."""
    cur = d
    for k in keys:
        if isinstance(cur, dict) and k in cur:
            cur = cur[k]
        else:
            return default
    return cur


def extract_summary(filepath):
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)

    row = {
        "Sample": os.path.splitext(os.path.basename(filepath))[0],
        "fastp_version": safe_get(data, "summary", "fastp_version"),
        "sequencing": safe_get(data, "summary", "sequencing"),

        # Before filtering
        "before_total_reads": safe_get(data, "summary", "before_filtering", "total_reads"),
        "before_total_bases": safe_get(data, "summary", "before_filtering", "total_bases"),
        "before_q20_rate": safe_get(data, "summary", "before_filtering", "q20_rate"),
        "before_q30_rate": safe_get(data, "summary", "before_filtering", "q30_rate"),
        "before_read1_mean_length": safe_get(data, "summary", "before_filtering", "read1_mean_length"),
        "before_read2_mean_length": safe_get(data, "summary", "before_filtering", "read2_mean_length"),
        "before_gc_content": safe_get(data, "summary", "before_filtering", "gc_content"),

        # After filtering
        "after_total_reads": safe_get(data, "summary", "after_filtering", "total_reads"),
        "after_total_bases": safe_get(data, "summary", "after_filtering", "total_bases"),
        "after_q20_rate": safe_get(data, "summary", "after_filtering", "q20_rate"),
        "after_q30_rate": safe_get(data, "summary", "after_filtering", "q30_rate"),
        "after_read1_mean_length": safe_get(data, "summary", "after_filtering", "read1_mean_length"),
        "after_read2_mean_length": safe_get(data, "summary", "after_filtering", "read2_mean_length"),
        "after_gc_content": safe_get(data, "summary", "after_filtering", "gc_content"),

        # Filtering result
        "passed_filter_reads": safe_get(data, "filtering_result", "passed_filter_reads"),
        "low_quality_reads": safe_get(data, "filtering_result", "low_quality_reads"),
        "too_many_N_reads": safe_get(data, "filtering_result", "too_many_N_reads"),
        "adapter_dimer_reads": safe_get(data, "filtering_result", "adapter_dimer_reads"),
        "too_short_reads": safe_get(data, "filtering_result", "too_short_reads"),
        "too_long_reads": safe_get(data, "filtering_result", "too_long_reads"),

        # Duplication / insert size
        "duplication_rate": safe_get(data, "duplication", "rate"),
        "insert_size_peak": safe_get(data, "insert_size", "peak"),

        # Adapter cutting
        "adapter_trimmed_reads": safe_get(data, "adapter_cutting", "adapter_trimmed_reads"),
        "adapter_trimmed_bases": safe_get(data, "adapter_cutting", "adapter_trimmed_bases"),
    }
    return row


# Natural sort so mixed sample numbering (1F, 2F, ..., 10F, 1L, 2L, ...) orders
# sensibly instead of plain alphabetical order (which would put "10F" before "2F").
def natural_key(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def main():
    if "..." in json_dir:
        sys.exit("Edit the paths at the top of this script (replace /.../.../ with your path).")

    pattern = os.path.join(json_dir, "*.json")
    json_files = sorted(glob.glob(pattern), key=natural_key)
    print(f"Found {len(json_files)} JSON file(s). Extracting...")

    rows = []
    for fp in json_files:
        try:
            rows.append(extract_summary(fp))
        except Exception as e:
            print(f"  ! Skipped {os.path.basename(fp)}: {e}")

    if not rows:
        sys.exit("No data extracted. Check that json_dir has valid fastp JSON reports.")

    rows.sort(key=lambda r: natural_key(r["Sample"]))

    headers = list(rows[0].keys())

    wb = Workbook()
    ws = wb.active
    ws.title = "fastp_summary"

    # Header row
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row[header])

    # Auto-width columns (approximate)
    for col_idx, header in enumerate(headers, start=1):
        max_len = max([len(str(header))] + [len(str(r[header])) for r in rows])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 30)

    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    wb.save(output_file)
    print(f"Done. Wrote {len(rows)} rows to {output_file}")


if __name__ == "__main__":
    main()
